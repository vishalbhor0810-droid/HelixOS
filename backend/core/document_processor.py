import os
import sys
import json
from typing import Dict, Any, List, Optional

try:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError
except ImportError:
    PdfReader = None
    PdfReadError = Exception

try:
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError
except ImportError:
    Document = None
    PackageNotFoundError = Exception

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:
    Presentation = None


def _result(status: str, path: str, text: str = "", error: Optional[str] = None) -> Dict[str, Any]:
    filename = os.path.basename(path)
    return {
        "status": status,
        "filename": filename,
        "path": os.path.abspath(path),
        "text": text,
        "textLength": len(text),
        "error": error
    }


import logging
from typing import Tuple

from backend.config import settings
from backend.core import vision_ocr

logger = logging.getLogger("helix.document_processor")

_ocr_reader = None


def _get_ocr_reader():
    """Lazy initialize EasyOCR reader."""
    global _ocr_reader
    if _ocr_reader is None:
        try:
            import easyocr
            _ocr_reader = easyocr.Reader(["en"], gpu=False, verbose=False)
        except Exception as e:
            logger.warning(f"EasyOCR initialization failed: {e}")
            _ocr_reader = None
    return _ocr_reader


def _ocr_with_confidence(img) -> Tuple[str, float]:
    """Run OCR on a PIL Image, numpy array, or file path. Returns (text, mean confidence 0-1),
    with confidence weighted by the length of each recognized fragment."""
    reader = _get_ocr_reader()
    if reader is None:
        return "", 0.0
    try:
        import numpy as np
        if hasattr(img, "convert"):
            img = np.array(img.convert("RGB"))
        results = reader.readtext(img, detail=1)
        texts = [text for _, text, _ in results]
        total_chars = sum(len(t) for t in texts)
        confidence = (
            sum(len(text) * conf for _, text, conf in results) / total_chars if total_chars else 0.0
        )
        return " ".join(texts).strip(), confidence
    except Exception as e:
        logger.debug(f"OCR extraction failed for image: {e}")
        return "", 0.0


def _extract_text_from_image(img) -> str:
    """Run OCR on a PIL Image, numpy array, or file path."""
    return _ocr_with_confidence(img)[0]


# Pages with less digital text than this are treated as scanned and sent to OCR
MIN_PAGE_TEXT_CHARS = 40


def _ocr_pdf_pages(path: str, reader, page_indices: List[int]) -> Dict[int, str]:
    """OCR the given 0-based page indices. Returns {page_index: text}.
    Pages EasyOCR reads with low confidence (typically handwriting) go to the vision model."""
    filename = os.path.basename(path)
    use_vision = vision_ocr.is_available()
    if _get_ocr_reader() is None and not use_vision:
        logger.warning(f"OCR unavailable; scanned pages in {filename} will be skipped")
        return {}

    ocr_text: Dict[int, str] = {}
    vision_pages = 0
    try:
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(path)
        try:
            for i in page_indices:
                page = pdf[i]
                try:
                    # Render page at 2x resolution for clean OCR accuracy
                    pil_image = page.render(scale=2.0).to_pil()
                    text, confidence = _ocr_with_confidence(pil_image)
                    ocr_text[i] = text
                    if use_vision and confidence < settings.VISION_MIN_CONFIDENCE:
                        if vision_pages >= settings.VISION_MAX_PAGES:
                            continue
                        vision_pages += 1
                        logger.info(
                            f"Page {i + 1} of {filename}: OCR confidence {confidence:.2f}, "
                            f"using vision model"
                        )
                        vision_text = vision_ocr.transcribe_page(pil_image)
                        if vision_text:
                            ocr_text[i] = vision_text
                        else:
                            logger.warning(
                                f"Vision transcription failed for page {i + 1} of {filename}; "
                                f"keeping low-confidence OCR text"
                            )
                finally:
                    page.close()
        finally:
            pdf.close()
        if vision_pages >= settings.VISION_MAX_PAGES:
            logger.info(f"Vision page limit ({settings.VISION_MAX_PAGES}) reached for {filename}")
    except Exception as ocr_err:
        logger.warning(f"pypdfium2 render failed for {path}, falling back to embedded images: {ocr_err}")
        # Fallback to embedded image extraction in pypdf
        from PIL import Image
        import io
        for i in page_indices:
            parts = []
            for img_file in getattr(reader.pages[i], "images", []):
                try:
                    img_text = _extract_text_from_image(Image.open(io.BytesIO(img_file.data)))
                    if img_text:
                        parts.append(img_text)
                except Exception:
                    continue
            ocr_text[i] = " ".join(parts).strip()
    return ocr_text


def _extract_pdf(path: str) -> Dict[str, Any]:
    if PdfReader is None:
        return _result("error", path, error="pypdf is not installed")
    try:
        reader = PdfReader(path)
        pages_text = []
        for page in reader.pages:
            try:
                pages_text.append((page.extract_text() or "").strip())
            except Exception:
                pages_text.append("")

        # OCR any page whose digital text is empty or too short (scanned / image-only)
        scanned_pages = [i for i, t in enumerate(pages_text) if len(t) < MIN_PAGE_TEXT_CHARS]
        if scanned_pages:
            logger.info(
                f"Running OCR on {len(scanned_pages)}/{len(pages_text)} pages of {os.path.basename(path)}"
            )
            for i, text in _ocr_pdf_pages(path, reader, scanned_pages).items():
                if len(text) > len(pages_text[i]):
                    pages_text[i] = text

        full_text = "\n\n".join(t for t in pages_text if t).strip()
        if not full_text:
            return _result("empty", path)
        return _result("success", path, text=full_text)
    except PdfReadError as e:
        return _result("error", path, error=f"PDF read error: {e}")
    except Exception as e:
        return _result("error", path, error=f"Failed to read PDF: {e}")


def _extract_docx(path: str) -> Dict[str, Any]:
    if Document is None:
        return _result("error", path, error="python-docx is not installed")
    try:
        doc = Document(path)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        
        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
                    
        full_text = "\n\n".join(paragraphs).strip()
        if not full_text:
            return _result("empty", path)
        return _result("success", path, text=full_text)
    except PackageNotFoundError as e:
        return _result("error", path, error=f"Invalid docx package: {e}")
    except Exception as e:
        return _result("error", path, error=f"Failed to read docx: {e}")


def _extract_txt(path: str) -> Dict[str, Any]:
    encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]
    for enc in encodings:
        try:
            with open(path, "r", encoding=enc) as f:
                text = f.read().strip()
                if not text:
                    return _result("empty", path)
                return _result("success", path, text=text)
        except UnicodeDecodeError:
            continue
        except Exception as e:
            return _result("error", path, error=f"Failed to read text file: {e}")
            
    return _result("error", path, error="Could not decode text file with standard encodings")


def _extract_image(path: str) -> Dict[str, Any]:
    """Extract text from standalone images (.png, .jpg, .jpeg, .webp, .bmp, .tiff) using OCR."""
    try:
        from PIL import Image
        with Image.open(path) as img:
            text = _extract_text_from_image(img)
            if not text:
                return _result("empty", path)
            return _result("success", path, text=text)
    except Exception as e:
        return _result("error", path, error=f"Failed to process image: {e}")


# Cap on rows read per sheet so a large data dump doesn't flood the index
MAX_ROWS_PER_SHEET = 5000


def _cell_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # 500.0 -> "500"
    return str(value).strip()


def _extract_xlsx(path: str) -> Dict[str, Any]:
    """Extract sheet contents, pairing each value with its column header for better retrieval."""
    if load_workbook is None:
        return _result("error", path, error="openpyxl is not installed")
    try:
        # read_only streams rows instead of loading the whole file;
        # data_only returns the last calculated value of a formula, not "=SUM(...)"
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            sections = []
            for ws in wb.worksheets:
                headers = None
                lines = []
                for row in ws.iter_rows(values_only=True):
                    cells = [_cell_str(v) for v in row]
                    if not any(cells):
                        continue
                    if headers is None:
                        headers = cells
                        lines.append(" | ".join(c for c in cells if c))
                        continue
                    pairs = []
                    for i, value in enumerate(cells):
                        if not value:
                            continue
                        header = headers[i] if i < len(headers) else ""
                        pairs.append(f"{header}: {value}" if header else value)
                    lines.append(" | ".join(pairs))
                    if len(lines) >= MAX_ROWS_PER_SHEET:
                        logger.info(f"Truncated sheet '{ws.title}' of {os.path.basename(path)} at {MAX_ROWS_PER_SHEET} rows")
                        break
                if lines:
                    sections.append(f"Sheet: {ws.title}\n" + "\n".join(lines))
        finally:
            wb.close()

        full_text = "\n\n".join(sections).strip()
        if not full_text:
            return _result("empty", path)
        return _result("success", path, text=full_text)
    except Exception as e:
        return _result("error", path, error=f"Failed to read Excel file: {e}")


def _pptx_shape_texts(shapes) -> List[str]:
    texts = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            texts.extend(_pptx_shape_texts(shape.shapes))
        elif shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text:
                texts.append(text)
        elif shape.has_table:
            for row in shape.table.rows:
                row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                if row_text:
                    texts.append(row_text)
    return texts


def _extract_pptx(path: str) -> Dict[str, Any]:
    """Extract slide text (text boxes, tables, grouped shapes) and speaker notes."""
    if Presentation is None:
        return _result("error", path, error="python-pptx is not installed")
    try:
        prs = Presentation(path)
        slides = []
        for number, slide in enumerate(prs.slides, 1):
            parts = _pptx_shape_texts(slide.shapes)
            if slide.has_notes_slide:
                notes_frame = slide.notes_slide.notes_text_frame
                notes = notes_frame.text.strip() if notes_frame is not None else ""
                if notes:
                    parts.append(f"Speaker notes: {notes}")
            if parts:
                slides.append(f"Slide {number}:\n" + "\n".join(parts))

        full_text = "\n\n".join(slides).strip()
        if not full_text:
            return _result("empty", path)
        return _result("success", path, text=full_text)
    except Exception as e:
        return _result("error", path, error=f"Failed to read PowerPoint file: {e}")


EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".xlsx": _extract_xlsx,
    ".xlsm": _extract_xlsx,
    ".pptx": _extract_pptx,
    ".txt": _extract_txt,
    ".md": _extract_txt,
    ".log": _extract_txt,
    ".csv": _extract_txt,
    ".png": _extract_image,
    ".jpg": _extract_image,
    ".jpeg": _extract_image,
    ".webp": _extract_image,
    ".bmp": _extract_image,
    ".tiff": _extract_image,
}


def extract_text(path: str) -> Dict[str, Any]:
    """Extract text from supported document formats."""
    if not os.path.exists(path):
        return _result("error", path, error="File does not exist")
        
    ext = os.path.splitext(path)[1].lower()
    extractor = EXTRACTORS.get(ext)
    if not extractor:
        return _result("unsupported", path, error=f"Unsupported file format: {ext}")
        
    return extractor(path)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(json.dumps({"status": "error", "error": "Missing file path argument"}))
        sys.exit(1)
        
    file_path = sys.argv[1]
    result = extract_text(file_path)
    print(json.dumps(result))
